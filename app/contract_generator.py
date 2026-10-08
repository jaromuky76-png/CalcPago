import os
import io
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def find_template_path():
    # 1. Search inside local templates folder (portable across machines and Render)
    cur = os.path.dirname(os.path.abspath(__file__))
    local_tmpl = os.path.join(cur, "templates", "ACTUALIZACION BORRADOR DE CONTRATO DE SERVICIOS DE INSTALACIÓN.docx")
    if os.path.exists(local_tmpl):
        return local_tmpl
    if os.path.isdir(os.path.join(cur, "templates")):
        for f in os.listdir(os.path.join(cur, "templates")):
            if "BORRADOR DE CONTRATO" in f and f.endswith(".docx"):
                return os.path.join(cur, "templates", f)

    # 2. Search upwards from current file to find the PROVEEDORES directory
    for _ in range(6):
        cand = os.path.join(cur, "PROVEEDORES", "ACTUALIZACION BORRADOR DE CONTRATO DE SERVICIOS DE INSTALACIÓN.docx")
        if os.path.exists(cand):
            return cand
        if os.path.isdir(os.path.join(cur, "PROVEEDORES")):
            for f in os.listdir(os.path.join(cur, "PROVEEDORES")):
                if "BORRADOR DE CONTRATO" in f and f.endswith(".docx"):
                    return os.path.join(cur, "PROVEEDORES", f)
        parent = os.path.dirname(cur)
        if parent == cur:
            break
        cur = parent
    return local_tmpl

DEFAULT_TEMPLATE_PATH = find_template_path()

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def number_to_words_es(n):
    unidades = ["cero", "un", "dos", "tres", "cuatro", "cinco", "seis", "siete", "ocho", "nueve",
                "diez", "once", "doce", "trece", "catorce", "quince", "dieciséis", "diecisiete",
                "dieciocho", "diecinueve", "veinte", "veintiuno", "veintidós", "veintitrés",
                "veinticuatro", "veinticinco", "veintiséis", "veintisiete", "veintiocho", "veintinueve", "treinta", "treinta y uno"]
    try:
        n = int(n)
        if 0 <= n < len(unidades):
            return unidades[n]
    except Exception:
        pass
    return str(n)

def monto_a_letras(amount):
    if amount is None or amount == 0:
        return 'CERO CON 00/100 CÓRDOBAS (C$ 0.00)'
    num = round(float(amount), 2)
    int_part = int(abs(num))
    cents = int(round((abs(num) - int_part) * 100))
    units = ['', 'UN', 'DOS', 'TRES', 'CUATRO', 'CINCO', 'SEIS', 'SIETE', 'OCHO', 'NUEVE']
    teens = ['DIEZ', 'ONCE', 'DOCE', 'TRECE', 'CATORCE', 'QUINCE', 'DIECISÉIS', 'DIECISIETE', 'DIECIOCHO', 'DIECINUEVE']
    tens = ['', 'DIEZ', 'VEINTE', 'TREINTA', 'CUARENTA', 'CINCUENTA', 'SESENTA', 'SETENTA', 'OCHENTA', 'NOVENTA']
    hundreds = ['', 'CIENTO', 'DOSCIENTOS', 'TRESCIENTOS', 'CUATROCIENTOS', 'QUINIENTOS', 'SEISCIENTOS', 'SETECIENTOS', 'OCHOCIENTOS', 'NOVECIENTOS']

    def section(n):
        if n == 100:
            return 'CIEN'
        h = n // 100
        rem = n % 100
        res = []
        if h > 0:
            res.append(hundreds[h])
        if rem > 0:
            if rem < 10:
                res.append(units[rem])
            elif rem < 20:
                res.append(teens[rem - 10])
            elif rem == 20:
                res.append('VEINTE')
            elif rem < 30:
                v_units = ['', 'VEINTIÚN', 'VEINTIDÓS', 'VEINTITRÉS', 'VEINTICUATRO', 'VEINTICINCO', 'VEINTISÉIS', 'VEINTISIETE', 'VEINTIOCHO', 'VEINTINUEVE']
                res.append(v_units[rem - 20])
            else:
                t = rem // 10
                u = rem % 10
                if u == 0:
                    res.append(tens[t])
                else:
                    res.append(tens[t] + ' Y ' + units[u])
        return ' '.join(res)

    if int_part == 0:
        words = 'CERO'
    else:
        parts = []
        millions = int_part // 1000000
        thousands = (int_part % 1000000) // 1000
        remainder = int_part % 1000
        if millions > 0:
            if millions == 1:
                parts.append('UN MILLÓN')
            else:
                parts.append(section(millions) + ' MILLONES')
        if thousands > 0:
            if thousands == 1:
                parts.append('MIL')
            else:
                parts.append(section(thousands) + ' MIL')
        if remainder > 0 or not parts:
            parts.append(section(remainder))
        words = ' '.join(parts).strip()

    cents_str = f'{cents:02d}'
    formatted_num = f'{num:,.2f}'
    return f'{words} CÓRDOBAS CON {cents_str}/100 CÓRDOBAS (C$ {formatted_num})'

def generate_contract_document(provider_data, template_path=None):
    if not template_path:
        template_path = DEFAULT_TEMPLATE_PATH
        
    doc = docx.Document(template_path)
    
    # Provider data unpacking
    nombre_rep = provider_data.get('nombre_representante', '').strip() or provider_data.get('nombre', '').strip() or '[PENDIENTE: REPRESENTANTE]'
    nombre_comercial = provider_data.get('nombre_comercial', '').strip() or provider_data.get('nombre', '').strip() or '[PENDIENTE: NOMBRE COMERCIAL]'
    cedula = provider_data.get('cedula', '').strip() or '[PENDIENTE: CÉDULA]'
    ruc = provider_data.get('ruc', '').strip()
    estado_civil = provider_data.get('estado_civil', 'mayor de edad').strip()
    profesion = provider_data.get('profesion', 'técnico').strip() or '[PENDIENTE: PROFESIÓN]'
    domicilio = provider_data.get('domicilio', 'Managua').strip()
    regimen = provider_data.get('regimen', 'Régimen de Cuota Fija').strip()
    
    banco = provider_data.get('banco', 'BAC Credomatic').strip()
    cuenta_bancaria = provider_data.get('cuenta_bancaria', '').strip() or '[PENDIENTE: CUENTA BANCARIA]'
    titular_cuenta = provider_data.get('titular_cuenta', '').strip() or nombre_rep
    
    direccion = provider_data.get('direccion', '').strip() or '[PENDIENTE: DIRECCIÓN]'
    telefono = provider_data.get('telefono', '').strip() or '[PENDIENTE: TELÉFONO]'
    correo = provider_data.get('correo', '').strip() or '[PENDIENTE: CORREO]'

    # Extended operational and contractual fields
    is_foraneo = (provider_data.get('contract_variant') == 'FORANEO')
    base_operativa = provider_data.get('base_operativa', 'Managua').strip() or 'Managua'
    departamentos = provider_data.get('departamentos', '').strip()
    geocerca_km = provider_data.get('geocerca_km', 14)
    condicion_pago = provider_data.get('condicion_pago', 'Semanal contra factura').strip()
    garantia_instalacion = provider_data.get('garantia_instalacion', 12)
    garantia_mantenimiento = provider_data.get('garantia_mantenimiento', '1')
    politica_uniformes = provider_data.get('politica_uniformes', 'Uniforme Oficial SINSA / Maestros').strip()
    contacto_operativo = provider_data.get('contacto_operativo', '').strip()
    enable_consignacion = provider_data.get('enable_consignacion', False) or provider_data.get('consignacion_activa', False)
    materiales = provider_data.get('materiales_consignados') or provider_data.get('materiales', [])
    
    # Date formatting
    dia = provider_data.get('dia', 23)
    mes = provider_data.get('mes', 'octubre')
    anio = provider_data.get('anio', 2026)
    dia_letras = number_to_words_es(dia)

    # 0. Proposal Banner for Foraneo Variant
    if is_foraneo and len(doc.paragraphs) > 0:
        p_banner = doc.paragraphs[0].insert_paragraph_before()
        p_banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_b1 = p_banner.add_run("PROPUESTA DE VARIANTE CONTRACTUAL FORÁNEA (v1.0-PROPUESTA)\n")
        r_b1.bold = True
        r_b1.font.name = 'Arial'
        r_b1.font.size = Pt(10.5)
        r_b1.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)
        r_b2 = p_banner.add_run(
            f"DOCUMENTO SUJETO A REVISIÓN LEGAL Y APROBACIÓN FORMAL DE GERENCIA GENERAL.\n"
            f"INCORPORA CUSTODIA DE INVENTARIO EN CONSIGNACIÓN Y GEOCERCA DEPARTAMENTAL ({base_operativa.upper()}).\n"
        )
        r_b2.italic = True
        r_b2.font.name = 'Arial'
        r_b2.font.size = Pt(9)
        r_b2.font.color.rgb = RGBColor(0x92, 0x40, 0x0E)
    
    # 1. Preamble (P2)
    if len(doc.paragraphs) > 2:
        p2 = doc.paragraphs[2]
        preamble_split = "y por otra parte,"
        if preamble_split in p2.text:
            sinsa_part = p2.text.split(preamble_split)[0] + preamble_split + " "
            contractor_part = (
                f"{nombre_rep.upper()}, mayor de edad, {estado_civil.lower()}, {profesion.lower()}, "
                f"con domicilio en {domicilio}, titular de cédula de identidad nicaragüense número: {cedula}"
            )
            if ruc:
                contractor_part += f" y cédula RUC: {ruc}"
                
            contractor_part += (
                f", quien actúa en nombre e interés de negocio bajo {regimen} denominado {nombre_comercial.upper()}, "
                f"quien en adelante se denominará EL CONTRATISTA, ambas partes de común acuerdo convenimos en celebrar el siguiente"
            )
            p2.text = sinsa_part + contractor_part
            for r in p2.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(10)

    # 1.1 Cláusula Primera: Objeto del contrato (Territorial coverage if foraneo)
    if is_foraneo:
        for p in doc.paragraphs:
            if "PRIMERA [OBJETO DEL CONTRATO]" in p.text:
                pass
            elif "requeridas por los CLIENTES del CONTRATANTE." in p.text:
                cobertura_str = f" en las circunscripciones y municipios autorizados del departamento de {base_operativa.upper()}"
                if departamentos and departamentos.upper() != base_operativa.upper():
                    cobertura_str += f" ({departamentos})"
                cobertura_str += "."
                p.text = p.text.replace("requeridas por los CLIENTES del CONTRATANTE.", f"requeridas por los CLIENTES del CONTRATANTE{cobertura_str}")
                for r in p.runs:
                    r.font.name = 'Arial'
                    r.font.size = Pt(10)
                break

    # 2. Cláusula Séptima (Forma de pago)
    for p in doc.paragraphs:
        if "Estos pagos serán realizados mediante transferencia bancaria" in p.text:
            if str(condicion_pago).startswith('CREDITO'):
                dias_cred = 'quince (15)' if '15' in str(condicion_pago) else 'treinta (30)'
                p.text = (
                    f"Estos pagos serán realizados bajo la modalidad de crédito comercial de {dias_cred} días calendario "
                    f"posteriores a la presentación formal de la factura, mediante transferencia bancaria a la cuenta de "
                    f"{banco.upper()} número: {cuenta_bancaria} en moneda córdobas a nombre de {titular_cuenta.upper()}."
                )
            else:
                p.text = f"Estos pagos serán realizados mediante transferencia bancaria a la cuenta de {banco.upper()} número: {cuenta_bancaria} en moneda córdobas a nombre de {titular_cuenta.upper()}."
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(10)
        elif "XXXXXXXXXXX en moneda córdobas a nombre del CONTRATISTA" in p.text:
            p.text = ""

    # 2.1 Garantías (DÉCIMA SEGUNDA)
    for idx, p in enumerate(doc.paragraphs):
        if "DÉCIMA SEGUNDA [GARANTÍA]" in p.text or "DECIMA SEGUNDA [GARANTIA]" in p.text:
            if idx + 1 < len(doc.paragraphs):
                p_gar = doc.paragraphs[idx + 1]
                p_gar.text = (
                    f"EL CONTRATISTA otorga una garantía de {garantia_instalacion} MESES calendario sobre la mano de obra "
                    f"de las instalaciones realizadas, contados a partir de la firma del Acta de Entrega y Recepción por el cliente. "
                    f"Si durante este plazo se presentaren fallas derivadas de una deficiente instalación, fuga de refrigerante por "
                    f"mala abocardadura o deficiencias en conexiones eléctricas, EL CONTRATISTA corregirá de inmediato el daño sin costo alguno. "
                    f"Asimismo, para los servicios de mantenimiento preventivo y limpieza técnica de equipos, EL CONTRATISTA otorga "
                    f"una garantía técnica de {garantia_mantenimiento} MES(ES) calendario. En todos los casos responderá ante cualquier "
                    f"reclamación o demanda por daños a terceros provocados en la ejecución de los servicios."
                )
                for r in p_gar.runs:
                    r.font.name = 'Arial'
                    r.font.size = Pt(10)
            break

    # 2.2 Política de Uniformes
    if politica_uniformes in ['EXCEPCION_PROPIA_GAFETE', 'PROVEEDOR_GAFETE']:
        for p in doc.paragraphs:
            if "portar en todo momento el uniforme reglamentario" in p.text or "prohíbe de manera expresa a EL CONTRATISTA y a su personal portar uniformes" in p.text:
                p.text += " De manera excepcional y sujeta a autorización escrita de Centro de Servicios, se autoriza a EL CONTRATISTA portar vestimenta técnica propia siempre que porte visiblemente el gafete o credencial oficial emitida por SINSA."
                for r in p.runs:
                    r.font.name = 'Arial'
                    r.font.size = Pt(10)
                break

    # 3. Cláusula Décima Séptima (Notificaciones)
    for p in doc.paragraphs:
        if "CONTRATISTA: XXXXXXXXXXX" in p.text:
            contact_op_str = f" / Contacto Operativo: {contacto_operativo}" if contacto_operativo else ""
            p.text = (
                f"CONTRATISTA: {nombre_comercial.upper()}, {direccion}. "
                f"Con Atención a: {nombre_rep}{contact_op_str}. Teléfono: {telefono}. "
                f"Correo Electrónico: {correo}."
            )
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(10)

    # 4. Cláusula Vigésima (Fecha de celebración)
    for p in doc.paragraphs:
        if "En fe de lo cual firmamos el presente contrato" in p.text:
            p.text = (
                f"En fe de lo cual firmamos el presente contrato, en dos tantos de un mismo tenor, "
                f"en la ciudad de Managua, a los {dia_letras} ({dia}) días del mes de {mes.lower()} del año {anio}."
            )
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(10)

    # 5. Signatures
    for p in doc.paragraphs:
        if "Eduin Jose Orozco Castro" in p.text or "Multiservicios Orozco" in p.text:
            p.text = p.text.replace("Eduin Jose Orozco Castro", nombre_rep).replace("Multiservicios Orozco", nombre_comercial)
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(10)

    # 5.1 Cláusula de Consignación en el cuerpo del contrato si aplica
    if enable_consignacion and len(materiales) > 0:
        subtotal_mat = sum(float(m.get('cantidad', 0) or 0) * float(m.get('costo_unitario', 0) or 0) for m in materiales)
        iva_mat = subtotal_mat * 0.15
        total_mat = subtotal_mat + iva_mat
        total_mat_letras = monto_a_letras(total_mat)

        p_cons = doc.add_paragraph()
        p_cons.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_c_title = p_cons.add_run("CLÁUSULA DE CONSIGNACIÓN DE MATERIALES E INVENTARIO INICIAL:\n")
        r_c_title.bold = True
        r_c_title.font.name = 'Arial'
        r_c_title.font.size = Pt(10)
        r_c_body = p_cons.add_run(
            f"Por medio de la presente cláusula, EL CONTRATANTE entrega a EL CONTRATISTA en calidad de consignación mercantil, "
            f"y éste recibe a su entera satisfacción en depósito responsable, el inventario inicial de materiales, accesorios y "
            f"repuestos para instalación de aires acondicionados detallado en el ANEXO II del presente contrato. "
            f"Las partes declaran expresamente que el valor total del inventario consignado asciende a la suma de {total_mat_letras}, "
            f"compuesto por un Subtotal Neto de C$ {subtotal_mat:,.2f} más el quince por ciento (15%) correspondiente al Impuesto al "
            f"Valor Agregado (IVA) por la suma de C$ {iva_mat:,.2f}. EL CONTRATISTA se constituye en custodio legal y depositario "
            f"mercantil de dichos bienes, asumiendo plena responsabilidad civil, comercial y penal por mermas injustificadas, pérdidas "
            f"o extravíos, obligándose a utilizarlos exclusivamente en la ejecución de las Órdenes de Trabajo encomendadas por "
            f"EL CONTRATANTE y a reponerlos o liquidarlos contra cada servicio facturado."
        )
        r_c_body.font.name = 'Arial'
        r_c_body.font.size = Pt(10)

    # 6. Build ANEXO I Table
    tarifas = provider_data.get('tarifas', [])
    tarifa_combustible = provider_data.get('tarifa_combustible', 12.0)
    
    # Add title for Annex table
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(f"ANEXO DE TARIFAS DE SERVICIOS - {nombre_comercial.upper()}")
    r_title.bold = True
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(11)
    r_title.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # Create Table: Rows = Header(2) + len(tarifas) + combustible(1), Cols = 3
    table = doc.add_table(rows=0, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(1.5), Inches(4.2), Inches(1.5)]

    # Main Banner Row
    row_banner = table.add_row()
    cell_banner = row_banner.cells[0]
    cell_banner.merge(row_banner.cells[1]).merge(row_banner.cells[2])
    set_cell_background(cell_banner, "1E293B")
    set_cell_margins(cell_banner, top=120, bottom=120, left=150, right=150)
    p_b = cell_banner.paragraphs[0]
    p_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_b = p_b.add_run("TABLA DE OFERTA CONTRATISTA DE MAESTROS")
    r_b.bold = True
    r_b.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r_b.font.name = 'Arial'
    r_b.font.size = Pt(11)

    # Column Headers Row
    headers = ["RMS", "DESCRIPCIÓN DE LA ACTIVIDAD", "TARIFA (C$)"]
    row_hdr = table.add_row()
    for idx, (cell, h_text) in enumerate(zip(row_hdr.cells, headers)):
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h_text)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)

    # Data Rows
    for idx, item in enumerate(tarifas):
        row = table.add_row()
        bg_color = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
        
        rms_val = str(item.get('rms', '') or '').strip()
        desc_val = str(item.get('descripcion', '') or '').strip()
        try:
            price_val = float(item.get('tarifa', 0) or 0)
        except Exception:
            price_val = 0.0
        
        # Col 0: RMS
        c0 = row.cells[0]
        set_cell_background(c0, bg_color)
        set_cell_margins(c0, top=80, bottom=80, left=120, right=120)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(rms_val if rms_val else "-")
        r0.font.name = 'Arial'
        r0.font.size = Pt(9)
        
        # Col 1: Desc
        c1 = row.cells[1]
        set_cell_background(c1, bg_color)
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(desc_val)
        r1.font.name = 'Arial'
        r1.font.size = Pt(9)
        
        # Col 2: Price
        c2 = row.cells[2]
        set_cell_background(c2, bg_color)
        set_cell_margins(c2, top=80, bottom=80, left=120, right=120)
        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r2 = p2.add_run(f"C$ {price_val:,.2f}")
        r2.bold = True
        r2.font.name = 'Arial'
        r2.font.size = Pt(9)

    # Combustible Row
    if tarifa_combustible is not None:
        row_fuel = table.add_row()
        c_f0 = row_fuel.cells[0]
        c_f1 = row_fuel.cells[1]
        c_f0.merge(c_f1)
        set_cell_background(c_f0, "F1F5F9")
        set_cell_margins(c_f0, top=90, bottom=90, left=120, right=120)
        pf = c_f0.paragraphs[0]
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        if is_foraneo:
            fuel_label = f"TARIFA DE COMBUSTIBLE POR KM FUERA DEL RADIO DE {base_operativa.upper()} ({geocerca_km} KM)"
        else:
            fuel_label = "TARIFA DE COMBUSTIBLE POR KM FUERA DEL RADIO DE LA CIUDAD (14 KM MANAGUA)"
            
        rf = pf.add_run(fuel_label)
        rf.bold = True
        rf.font.name = 'Arial'
        rf.font.size = Pt(9)
        
        c_f2 = row_fuel.cells[2]
        set_cell_background(c_f2, "F1F5F9")
        set_cell_margins(c_f2, top=90, bottom=90, left=120, right=120)
        pf2 = c_f2.paragraphs[0]
        pf2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        try:
            fuel_val = float(tarifa_combustible)
        except Exception:
            fuel_val = 0.0
        rf2 = pf2.add_run(f"C$ {fuel_val:,.2f}")
        rf2.bold = True
        rf2.font.name = 'Arial'
        rf2.font.size = Pt(9)

    # Set cell widths
    for row in table.rows:
        for i, w in enumerate(col_widths):
            if i < len(row.cells):
                row.cells[i].width = w

    # 7. ANEXO II: TABLA DE MATERIALES EN CONSIGNACIÓN
    if enable_consignacion and len(materiales) > 0:
        p_anx2_title = doc.add_paragraph()
        p_anx2_title.paragraph_format.page_break_before = True
        p_anx2_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2_title = p_anx2_title.add_run("ANEXO II: TABLA DE MATERIALES E INVENTARIO INICIAL EN CONSIGNACIÓN MERCANTIL")
        r2_title.bold = True
        r2_title.font.name = 'Arial'
        r2_title.font.size = Pt(11)
        r2_title.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

        p_anx2_sub = doc.add_paragraph()
        p_anx2_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2_sub = p_anx2_sub.add_run(f"CONTRATISTA: {nombre_comercial.upper()} | BASE OPERATIVA: {base_operativa.upper()}")
        r2_sub.font.name = 'Arial'
        r2_sub.font.size = Pt(9.5)
        r2_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

        table_mat = doc.add_table(rows=0, cols=6)
        table_mat.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_mat.autofit = False

        col_mat_widths = [Inches(1.0), Inches(2.6), Inches(0.7), Inches(0.8), Inches(1.0), Inches(1.1)]

        # Header Row
        mat_headers = ["CÓDIGO RMS", "DESCRIPCIÓN DEL INSUMO", "UNIDAD", "CANTIDAD", "COSTO UNIT.", "SUBTOTAL"]
        row_mat_hdr = table_mat.add_row()
        for idx, (cell, h_text) in enumerate(zip(row_mat_hdr.cells, mat_headers)):
            set_cell_background(cell, "0F172A")
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in [0, 2] else (WD_ALIGN_PARAGRAPH.RIGHT if idx >= 3 else WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(h_text)
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r.font.name = 'Arial'
            r.font.size = Pt(9)

        calc_subtotal = 0.0
        for idx, m in enumerate(materiales):
            row_m = table_mat.add_row()
            bg_color = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
            
            c_rms = str(m.get('rms', '') or '').strip()
            c_desc = str(m.get('descripcion', '') or '').strip()
            c_und = str(m.get('unidad', 'UND') or 'UND').strip()
            try:
                c_qty = float(m.get('cantidad', 0) or 0)
            except Exception:
                c_qty = 0.0
            try:
                c_cost = float(m.get('costo_unitario', 0) or 0)
            except Exception:
                c_cost = 0.0
            c_sub = c_qty * c_cost
            calc_subtotal += c_sub

            vals = [
                (c_rms or "-", WD_ALIGN_PARAGRAPH.CENTER, False),
                (c_desc, WD_ALIGN_PARAGRAPH.LEFT, False),
                (c_und, WD_ALIGN_PARAGRAPH.CENTER, False),
                (f"{c_qty:g}", WD_ALIGN_PARAGRAPH.RIGHT, False),
                (f"C$ {c_cost:,.2f}", WD_ALIGN_PARAGRAPH.RIGHT, False),
                (f"C$ {c_sub:,.2f}", WD_ALIGN_PARAGRAPH.RIGHT, True)
            ]
            for c_idx, (text_val, align_val, bold_val) in enumerate(vals):
                cell_curr = row_m.cells[c_idx]
                set_cell_background(cell_curr, bg_color)
                set_cell_margins(cell_curr, top=70, bottom=70, left=80, right=80)
                p_c = cell_curr.paragraphs[0]
                p_c.alignment = align_val
                r_c = p_c.add_run(text_val)
                r_c.bold = bold_val
                r_c.font.name = 'Arial'
                r_c.font.size = Pt(8.5)

        calc_iva = calc_subtotal * 0.15
        calc_total = calc_subtotal + calc_iva

        # Subtotal Row
        row_sub = table_mat.add_row()
        c_sub0 = row_sub.cells[0]
        for merge_idx in range(1, 5):
            c_sub0.merge(row_sub.cells[merge_idx])
        set_cell_background(c_sub0, "F8FAFC")
        set_cell_margins(c_sub0, top=80, bottom=80, left=100, right=100)
        p_s0 = c_sub0.paragraphs[0]
        p_s0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_s0 = p_s0.add_run("SUBTOTAL NETO DE MATERIALES:")
        r_s0.bold = True
        r_s0.font.name = 'Arial'
        r_s0.font.size = Pt(9)

        c_sub_val = row_sub.cells[5]
        set_cell_background(c_sub_val, "F8FAFC")
        set_cell_margins(c_sub_val, top=80, bottom=80, left=100, right=100)
        p_sv = c_sub_val.paragraphs[0]
        p_sv.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_sv = p_sv.add_run(f"C$ {calc_subtotal:,.2f}")
        r_sv.bold = True
        r_sv.font.name = 'Arial'
        r_sv.font.size = Pt(9)

        # IVA Row
        row_iva = table_mat.add_row()
        c_iva0 = row_iva.cells[0]
        for merge_idx in range(1, 5):
            c_iva0.merge(row_iva.cells[merge_idx])
        set_cell_background(c_iva0, "F8FAFC")
        set_cell_margins(c_iva0, top=80, bottom=80, left=100, right=100)
        p_i0 = c_iva0.paragraphs[0]
        p_i0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_i0 = p_i0.add_run("IMPUESTO AL VALOR AGREGADO (IVA 15%):")
        r_i0.bold = True
        r_i0.font.name = 'Arial'
        r_i0.font.size = Pt(9)

        c_iva_val = row_iva.cells[5]
        set_cell_background(c_iva_val, "F8FAFC")
        set_cell_margins(c_iva_val, top=80, bottom=80, left=100, right=100)
        p_iv = c_iva_val.paragraphs[0]
        p_iv.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_iv = p_iv.add_run(f"C$ {calc_iva:,.2f}")
        r_iv.bold = True
        r_iv.font.name = 'Arial'
        r_iv.font.size = Pt(9)

        # Total Row
        row_tot = table_mat.add_row()
        c_tot0 = row_tot.cells[0]
        for merge_idx in range(1, 5):
            c_tot0.merge(row_tot.cells[merge_idx])
        set_cell_background(c_tot0, "EFF6FF")
        set_cell_margins(c_tot0, top=90, bottom=90, left=100, right=100)
        p_t0 = c_tot0.paragraphs[0]
        p_t0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_t0 = p_t0.add_run("TOTAL VALORIZADO EN CONSIGNACIÓN:")
        r_t0.bold = True
        r_t0.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
        r_t0.font.name = 'Arial'
        r_t0.font.size = Pt(9.5)

        c_tot_val = row_tot.cells[5]
        set_cell_background(c_tot_val, "EFF6FF")
        set_cell_margins(c_tot_val, top=90, bottom=90, left=100, right=100)
        p_tv = c_tot_val.paragraphs[0]
        p_tv.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_tv = p_tv.add_run(f"C$ {calc_total:,.2f}")
        r_tv.bold = True
        r_tv.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
        r_tv.font.name = 'Arial'
        r_tv.font.size = Pt(9.5)

        # Set cell widths
        for r in table_mat.rows:
            for i, w in enumerate(col_mat_widths):
                if i < len(r.cells):
                    r.cells[i].width = w

        # Words Note
        p_note = doc.add_paragraph()
        p_note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_n1 = p_note.add_run("VALOR TOTAL EN LETRAS: ")
        r_n1.bold = True
        r_n1.font.name = 'Arial'
        r_n1.font.size = Pt(9)
        r_n2 = p_note.add_run(monto_a_letras(calc_total))
        r_n2.font.name = 'Arial'
        r_n2.font.size = Pt(9)

    target_stream = io.BytesIO()
    doc.save(target_stream)
    target_stream.seek(0)
    return target_stream
