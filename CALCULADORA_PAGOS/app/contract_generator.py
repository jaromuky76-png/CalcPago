import os
import io
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def find_template_path():
    # 1. Search inside local templates folder (portable across machines and Render)
    cur = os.path.dirname(os.path.abspath(__file__))
    local_tmpl = os.path.join(cur, "templates", "ACTUALIZACION BORRADOR DE CONTRATO DE SERVICIOS DE INSTALACIÓN.docx")
    if os.path.exists(local_tmpl):
        return local_tmpl
    if os.path.isdir(os.path.join(cur, "templates")):
        for f in os.listdir(os.path.join(cur, "templates")):
            if "BORRADOR DE CONTRATO" in f and f.endswith(".docx"):
                return os.path.join(cur, "templates", f)

    # 2. Search upwards from current file to find the PROVEEDORES directory
    for _ in range(6):
        cand = os.path.join(cur, "PROVEEDORES", "ACTUALIZACION BORRADOR DE CONTRATO DE SERVICIOS DE INSTALACIÓN.docx")
        if os.path.exists(cand):
            return cand
        if os.path.isdir(os.path.join(cur, "PROVEEDORES")):
            for f in os.listdir(os.path.join(cur, "PROVEEDORES")):
                if "BORRADOR DE CONTRATO" in f and f.endswith(".docx"):
                    return os.path.join(cur, "PROVEEDORES", f)
        parent = os.path.dirname(cur)
        if parent == cur:
            break
        cur = parent
    return local_tmpl

DEFAULT_TEMPLATE_PATH = find_template_path()

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def number_to_words_es(n):
    unidades = ["cero", "un", "dos", "tres", "cuatro", "cinco", "seis", "siete", "ocho", "nueve",
                "diez", "once", "doce", "trece", "catorce", "quince", "dieciséis", "diecisiete",
                "dieciocho", "diecinueve", "veinte", "veintiuno", "veintidós", "veintitrés",
                "veinticuatro", "veinticinco", "veintiséis", "veintisiete", "veintiocho", "veintinueve", "treinta", "treinta y uno"]
    try:
        n = int(n)
        if 0 <= n < len(unidades):
            return unidades[n]
    except Exception:
        pass
    return str(n)

def generate_contract_document(provider_data, template_path=None):
    if not template_path:
        template_path = DEFAULT_TEMPLATE_PATH
        
    doc = docx.Document(template_path)
    
    # Provider data unpacking
    nombre_rep = provider_data.get('nombre_representante', '').strip() or provider_data.get('nombre', '').strip() or '[PENDIENTE: REPRESENTANTE]'
    nombre_comercial = provider_data.get('nombre_comercial', '').strip() or provider_data.get('nombre', '').strip() or '[PENDIENTE: NOMBRE COMERCIAL]'
    cedula = provider_data.get('cedula', '').strip() or '[PENDIENTE: CÉDULA]'
    ruc = provider_data.get('ruc', '').strip()
    estado_civil = provider_data.get('estado_civil', 'mayor de edad').strip()
    profesion = provider_data.get('profesion', 'técnico').strip() or '[PENDIENTE: PROFESIÓN]'
    domicilio = provider_data.get('domicilio', 'Managua').strip()
    regimen = provider_data.get('regimen', 'Régimen de Cuota Fija').strip()
    
    banco = provider_data.get('banco', 'BAC Credomatic').strip()
    cuenta_bancaria = provider_data.get('cuenta_bancaria', '').strip() or '[PENDIENTE: CUENTA BANCARIA]'
    titular_cuenta = provider_data.get('titular_cuenta', '').strip() or nombre_rep
    
    direccion = provider_data.get('direccion', '').strip() or '[PENDIENTE: DIRECCIÓN]'
    telefono = provider_data.get('telefono', '').strip() or '[PENDIENTE: TELÉFONO]'
    correo = provider_data.get('correo', '').strip() or '[PENDIENTE: CORREO]'
    
    # Date formatting
    dia = provider_data.get('dia', 23)
    mes = provider_data.get('mes', 'octubre')
    anio = provider_data.get('anio', 2026)
    dia_letras = number_to_words_es(dia)
    
    # 1. Preamble (P2)
    p2 = doc.paragraphs[2]
    preamble_split = "y por otra parte,"
    if preamble_split in p2.text:
        sinsa_part = p2.text.split(preamble_split)[0] + preamble_split + " "
        contractor_part = (
            f"{nombre_rep.upper()}, mayor de edad, {estado_civil.lower()}, {profesion.lower()}, "
            f"con domicilio en {domicilio}, titular de cédula de identidad nicaragüense número: {cedula}"
        )
        if ruc:
            contractor_part += f" y cédula RUC: {ruc}"
            
        contractor_part += (
            f", quien actúa en nombre e interés de negocio bajo {regimen} denominado {nombre_comercial.upper()}, "
            f"quien en adelante se denominará EL CONTRATISTA, ambas partes de común acuerdo convenimos en celebrar el siguiente"
        )
        p2.text = sinsa_part + contractor_part
        for r in p2.runs:
            r.font.name = 'Arial'
            r.font.size = Pt(10)

    # 2. Cláusula Séptima (P50, P51)
    for p in doc.paragraphs:
        if "Estos pagos serán realizados mediante transferencia bancaria" in p.text:
            p.text = f"Estos pagos serán realizados mediante transferencia bancaria a la cuenta de {banco.upper()} número: {cuenta_bancaria} en moneda córdobas a nombre de {titular_cuenta.upper()}."
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(10)
        elif "XXXXXXXXXXX en moneda córdobas a nombre del CONTRATISTA" in p.text:
            p.text = ""

    # 3. Cláusula Décima Séptima (Notificaciones)
    for p in doc.paragraphs:
        if "CONTRATISTA: XXXXXXXXXXX" in p.text:
            p.text = (
                f"CONTRATISTA: {nombre_comercial.upper()}, {direccion}. "
                f"Con Atención a: {nombre_rep}. Teléfono: {telefono}. "
                f"Correo Electrónico: {correo}."
            )
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(10)

    # 4. Cláusula Vigésima (Fecha de celebración)
    for p in doc.paragraphs:
        if "En fe de lo cual firmamos el presente contrato" in p.text:
            p.text = (
                f"En fe de lo cual firmamos el presente contrato, en dos tantos de un mismo tenor, "
                f"en la ciudad de Managua, a los {dia_letras} ({dia}) días del mes de {mes.lower()} del año {anio}."
            )
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(10)

    # 5. Signatures
    for p in doc.paragraphs:
        if "Eduin Jose Orozco Castro" in p.text or "Multiservicios Orozco" in p.text:
            p.text = p.text.replace("Eduin Jose Orozco Castro", nombre_rep).replace("Multiservicios Orozco", nombre_comercial)
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(10)

    # 6. Build ANEXO I Table
    tarifas = provider_data.get('tarifas', [])
    tarifa_combustible = provider_data.get('tarifa_combustible', 12.0)
    
    # Add title for Annex table
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(f"ANEXO DE TARIFAS DE SERVICIOS - {nombre_comercial.upper()}")
    r_title.bold = True
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(11)
    r_title.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # Create Table: Rows = Header(2) + len(tarifas) + combustible(1), Cols = 3
    table = doc.add_table(rows=0, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(1.5), Inches(4.2), Inches(1.5)]

    # Main Banner Row
    row_banner = table.add_row()
    cell_banner = row_banner.cells[0]
    cell_banner.merge(row_banner.cells[1]).merge(row_banner.cells[2])
    set_cell_background(cell_banner, "1E293B")
    set_cell_margins(cell_banner, top=120, bottom=120, left=150, right=150)
    p_b = cell_banner.paragraphs[0]
    p_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_b = p_b.add_run("TABLA DE OFERTA CONTRATISTA DE MAESTROS")
    r_b.bold = True
    r_b.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r_b.font.name = 'Arial'
    r_b.font.size = Pt(11)

    # Column Headers Row
    headers = ["RMS", "DESCRIPCIÓN DE LA ACTIVIDAD", "TARIFA (C$)"]
    row_hdr = table.add_row()
    for idx, (cell, h_text) in enumerate(zip(row_hdr.cells, headers)):
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h_text)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)

    # Data Rows
    for idx, item in enumerate(tarifas):
        row = table.add_row()
        bg_color = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
        
        rms_val = str(item.get('rms', '') or '').strip()
        desc_val = str(item.get('descripcion', '') or '').strip()
        try:
            price_val = float(item.get('tarifa', 0) or 0)
        except Exception:
            price_val = 0.0
        
        # Col 0: RMS
        c0 = row.cells[0]
        set_cell_background(c0, bg_color)
        set_cell_margins(c0, top=80, bottom=80, left=120, right=120)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(rms_val if rms_val else "-")
        r0.font.name = 'Arial'
        r0.font.size = Pt(9)
        
        # Col 1: Desc
        c1 = row.cells[1]
        set_cell_background(c1, bg_color)
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(desc_val)
        r1.font.name = 'Arial'
        r1.font.size = Pt(9)
        
        # Col 2: Price
        c2 = row.cells[2]
        set_cell_background(c2, bg_color)
        set_cell_margins(c2, top=80, bottom=80, left=120, right=120)
        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r2 = p2.add_run(f"C$ {price_val:,.2f}")
        r2.bold = True
        r2.font.name = 'Arial'
        r2.font.size = Pt(9)

    # Combustible Row
    if tarifa_combustible is not None:
        row_fuel = table.add_row()
        c_f0 = row_fuel.cells[0]
        c_f1 = row_fuel.cells[1]
        c_f0.merge(c_f1)
        set_cell_background(c_f0, "F1F5F9")
        set_cell_margins(c_f0, top=90, bottom=90, left=120, right=120)
        pf = c_f0.paragraphs[0]
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        rf = pf.add_run("TARIFA DE COMBUSTIBLE POR KM FUERA DEL RADIO DE LA CIUDAD")
        rf.bold = True
        rf.font.name = 'Arial'
        rf.font.size = Pt(9)
        
        c_f2 = row_fuel.cells[2]
        set_cell_background(c_f2, "F1F5F9")
        set_cell_margins(c_f2, top=90, bottom=90, left=120, right=120)
        pf2 = c_f2.paragraphs[0]
        pf2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        try:
            fuel_val = float(tarifa_combustible)
        except Exception:
            fuel_val = 0.0
        rf2 = pf2.add_run(f"C$ {fuel_val:,.2f}")
        rf2.bold = True
        rf2.font.name = 'Arial'
        rf2.font.size = Pt(9)

    # Set cell widths
    for row in table.rows:
        for i, w in enumerate(col_widths):
            if i < len(row.cells):
                row.cells[i].width = w

    target_stream = io.BytesIO()
    doc.save(target_stream)
    target_stream.seek(0)
    return target_stream
